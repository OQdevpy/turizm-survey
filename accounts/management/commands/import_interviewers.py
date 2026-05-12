"""Excel Sheet2 dan 165 ta intervyuverni o'qib, User + StaffProfile yaratadi.

Sheet2 strukturasi (3 ustun):
- Hudud (region code, 4 raqamli)
- Intervyuverlar (F.I.SH)
- post (6 raqamli post kodi — har bir xodim qaysi postga biriktirilgan)
"""
import re
import unicodedata
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import User
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from accounts.models import PostalOffice, StaffProfile


def slugify_username(full_name: str) -> str:
    """F.I.SH dan username yaratish: 'Tajibekov Baxtiyar Nietullaevich' -> 'tajibekov_baxtiyar'."""
    parts = full_name.strip().split()
    if len(parts) < 2:
        base = parts[0] if parts else 'user'
    else:
        base = f"{parts[0]}_{parts[1]}"
    base = unicodedata.normalize('NFKD', base).encode('ascii', 'ignore').decode('ascii')
    base = re.sub(r'[^a-zA-Z0-9_]+', '_', base).lower().strip('_')
    return base or 'user'


def split_name(full_name: str):
    """'Tajibekov Baxtiyar Nietullaevich' -> ('Baxtiyar', 'Tajibekov Nietullaevich')."""
    parts = full_name.strip().split()
    if len(parts) == 0:
        return ('', '')
    if len(parts) == 1:
        return (parts[0], '')
    if len(parts) == 2:
        return (parts[1], parts[0])
    first_name = parts[1]
    last_name = parts[0] + ' ' + ' '.join(parts[2:])
    return (first_name, last_name[:150])


class Command(BaseCommand):
    help = "Excel'dan 165 ta intervyuverni import qiladi (Sheet2 'post' ustuni asosida)"

    def add_arguments(self, parser):
        parser.add_argument(
            '--xlsx',
            default=str(Path(settings.BASE_DIR).parent / 'tz' / 'Postlar va intervyuverlar+.xlsx'),
            help='Excel fayl yo\'li',
        )
        parser.add_argument(
            '--password',
            default='staff2026',
            help='Barcha xodimlarga standart parol (default: staff2026)',
        )
        parser.add_argument(
            '--reset',
            action='store_true',
            help='Avvalgi xodimlarni o\'chirib qayta yaratish (superuser saqlab qolinadi)',
        )

    @transaction.atomic
    def handle(self, *args, **options):
        xlsx_path = Path(options['xlsx'])
        if not xlsx_path.exists():
            raise CommandError(f"Excel fayl topilmadi: {xlsx_path}")

        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError("openpyxl o'rnatilmagan: pip install openpyxl")

        password = options['password']

        # Barcha postlarni xaritada saqlash
        offices_by_code = {p.code: p for p in PostalOffice.objects.all()}
        if not offices_by_code:
            raise CommandError(
                "Postlar bazasida yo'q. Avval 'python manage.py seed_data' ishga tushiring."
            )

        self.stdout.write(f"Postlar mavjud: {len(offices_by_code)} ta")

        if options.get('reset'):
            non_super = StaffProfile.objects.exclude(user__is_superuser=True)
            count = non_super.count()
            user_ids = list(non_super.values_list('user_id', flat=True))
            non_super.delete()
            User.objects.filter(id__in=user_ids, is_superuser=False).delete()
            self.stdout.write(self.style.WARNING(f"[!] {count} ta eski xodim o'chirildi"))

        # Excel o'qish
        wb = load_workbook(xlsx_path, data_only=True)
        try:
            ws = wb['Sheet2']
        except KeyError:
            raise CommandError("Excel'da 'Sheet2' topilmadi")

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise CommandError("Sheet2 bo'sh")

        header = rows[0]
        self.stdout.write(f"Sheet2 ustunlari: {header}")
        if len(header) < 3:
            raise CommandError(
                f"Sheet2 da 3 ustun kutilgan (Hudud, Intervyuverlar, post), "
                f"lekin {len(header)} ta topildi"
            )

        data_rows = rows[1:]
        created = updated = skipped = 0
        username_seq = {}

        for row in data_rows:
            if not row or all(v is None for v in row):
                continue
            region_raw, name_raw, post_raw = row[0], row[1], row[2]
            if not name_raw:
                continue
            region_code = str(region_raw).strip() if region_raw is not None else ''
            full_name = str(name_raw).strip()
            post_code = str(post_raw).strip() if post_raw is not None else ''

            if not post_code:
                self.stdout.write(self.style.WARNING(
                    f"[SKIP] '{full_name}' uchun post kodi bo'sh"
                ))
                skipped += 1
                continue

            office = offices_by_code.get(post_code)
            if not office:
                self.stdout.write(self.style.WARNING(
                    f"[SKIP] '{full_name}' uchun post '{post_code}' bazada yo'q"
                ))
                skipped += 1
                continue

            base_username = slugify_username(full_name)
            # Unique username yaratish
            if base_username not in username_seq:
                username_seq[base_username] = 0
                username = base_username
            else:
                username_seq[base_username] += 1
                username = f"{base_username}{username_seq[base_username]}"
            # Agar yana band bo'lsa — keyingi nomerga o'tish
            while User.objects.filter(username=username).exclude(
                    staff_profile__full_name=full_name).exists():
                username_seq[base_username] += 1
                username = f"{base_username}{username_seq[base_username]}"

            first_name, last_name = split_name(full_name)

            user, was_created = User.objects.get_or_create(
                username=username,
                defaults={
                    'first_name': first_name,
                    'last_name': last_name,
                    'is_staff': True,
                    'is_active': True,
                },
            )
            if was_created:
                user.set_password(password)
                user.save()
                created += 1
            else:
                user.first_name = first_name
                user.last_name = last_name
                user.is_staff = True
                user.is_active = True
                user.set_password(password)
                user.save()
                updated += 1

            StaffProfile.objects.update_or_create(
                user=user,
                defaults={
                    'postal_office': office,
                    'full_name': full_name,
                    'region_code': region_code,
                    'position': 'Intervyuver',
                },
            )

        total = User.objects.filter(is_staff=True, is_superuser=False).count()
        self.stdout.write(self.style.SUCCESS(
            f"\n[OK] {created} ta yangi, {updated} ta yangilangan, {skipped} ta o'tkazib yuborilgan"
        ))
        self.stdout.write(self.style.SUCCESS(
            f"Jami faol xodimlar: {total}\n"
            f"Standart parol: '{password}'\n"
        ))

        # Post bo'yicha taqsimot
        from django.db.models import Count
        self.stdout.write("\nPost bo'yicha taqsimot:")
        for r in (StaffProfile.objects.values('postal_office__code', 'postal_office__name')
                  .annotate(c=Count('id')).order_by('postal_office__code')):
            name = (r['postal_office__name'] or '')[:55]
            self.stdout.write(f"  {r['postal_office__code']}: {r['c']:2d} ta — {name}")
