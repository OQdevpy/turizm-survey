"""Master sheet (barcha savollar bitta sheetda) — namuna XLSX yaratuvchi script.

Bu Django'siz ishlaydigan standalone script. openpyxl talab qiladi.

Foydalanish:
    python scripts/generate_master_sheet_sample.py

Natija: scripts/master_sheet_sample.xlsx
"""
import sys
from pathlib import Path
from datetime import datetime, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl o'rnatilmagan. `pip install openpyxl` ishga tushiring.")
    sys.exit(1)


# ============================================================
# Ustunlar tuzilishi
# ============================================================

# Inbound shaharlar
INBOUND_CITIES = [
    'Tashkent', 'Samarkand', 'Bukhara', 'Khiva', 'Shakhrisabz', 'Termez',
    'Kokand', 'Fergana', 'Namangan', 'Andijan', 'Nukus', 'Urgench', 'Other',
]
# Inbound ovqat joylari
INBOUND_FOOD = [
    'restaurants', 'street', 'fastfood', 'national',
    'friends_home', 'rented', 'other',
]
# Inbound Q18 reytinglar (12 ta)
INBOUND_RATINGS = [
    'r0_intl_transport', 'r1_passport_control', 'r2_hospitality',
    'r3_value_for_money', 'r4_food', 'r5_cleanliness',
    'r6_local_transport', 'r7_safety', 'r8_culture',
    'r9_accommodation', 'r10_health', 'r11_communication',
]
# Outbound Q14 xarajat qatorlari (13 ta + 1.1, 1.2 sub-rows)
OUTBOUND_EXP_ROWS = [
    '1', '1.1', '1.2', '2', '3', '4', '5', '6', '7', '8', '9', '10', '11', '12', '13',
]


def build_columns():
    """Barcha ustunlarni list qilib qaytaradi."""
    cols = []

    # 1) Identifikatsiya
    cols += [
        '#', 'ID', 'Turi', 'Manba', 'Til',
        'Boshlangan', 'Yakunlangan', "To'ldirish (s)",
    ]
    # 2) Xodim va bo'lim
    cols += [
        'Xodim login', 'Xodim F.I.SH', 'Xodim telefon',
        "Bo'lim kodi", "Bo'lim nomi", 'Region kodi', 'Post turi',
    ]
    # 3) Tarmoq va device
    cols += [
        'IP', 'Subnet /24', 'Device turi', 'OS', 'Brauzer',
        'Screen', 'Platform', 'Til (client)', 'Timezone', 'User-Agent',
    ]
    # 4) GPS
    cols += [
        'GPS bormi', 'Latitude', 'Longitude', 'GPS aniqlik (m)', 'Google Maps URL',
    ]
    # 5) Screening
    cols += ['Filtr holati', 'F1', 'F2', 'F3']
    # 6) Auto-summary
    cols += ['Davlat', 'Maqsad', 'Tunlar', 'Xarajat (umumiy)', 'Valyuta']

    # 7) INBOUND
    cols += [
        'IN.Q1 (Doimiy davlat)', 'IN.Q2 (Pasport)', 'IN.Q2_country (Boshqa pasport)',
        'IN.Q3 (Maqsad)', 'IN.Q4 (Biznes turi)',
        'IN.Q5_nights', 'IN.Q5_zero',
    ]
    cols += [f'IN.Q6.{c}' for c in INBOUND_CITIES]
    cols += ['IN.Q7 (Turar joy)']
    cols += [f'IN.Q8.{f}' for f in INBOUND_FOOD]
    cols += [
        'IN.Q9 (Paket?)',
        'IN.Q10_total', 'IN.Q10_uz', 'IN.Q11 (Paket kishi)',
        'IN.Q12_amount', 'IN.Q12_currency',
        'IN.Q13 (Kelish)', 'IN.Q13_airline',
        'IN.Q14 (Ketish)', 'IN.Q14_airline',
        'IN.Q15 (Daromad ulushi)',
        'IN.Q16_sum', 'IN.Q16_currency', 'IN.Q16_persons',
    ]
    # Q17 xarajatlar (14 qator × 3)
    for n in range(1, 15):
        cols += [f'IN.Q17.r{n}.amount', f'IN.Q17.r{n}.currency', f'IN.Q17.r{n}.inPackage']
    # Q18 reytinglar
    cols += [f'IN.Q18.{r}' for r in INBOUND_RATINGS]
    cols += ['IN.Q19 (Izoh)']

    # 8) OUTBOUND
    cols += [
        'OUT.Q1 (Asosiy davlat)', 'OUT.Q2 (Maqsad)', 'OUT.Q3 (Biznes turi)',
        'OUT.Q4_val (Tunlar)',
        'OUT.Q5 (Turar joy)', 'OUT.Q6 (Paket?)',
        'OUT.Q7 (Paket tunlari)', 'OUT.Q8 (Paket kishi)',
        'OUT.Q9_amount', 'OUT.Q9_currency',
        'OUT.Q10 (Chiqish)', 'OUT.Q10_airline',
        'OUT.Q11 (Qaytish)', 'OUT.Q11_airline',
        'OUT.Q12 (Daromad ulushi)',
        'OUT.Q13_amount', 'OUT.Q13_currency', 'OUT.Q13_persons',
    ]
    # Q14 xarajatlar (13+2 qator × 3)
    for n in OUTBOUND_EXP_ROWS:
        cols += [f'OUT.Q14.r{n}.amount', f'OUT.Q14.r{n}.currency', f'OUT.Q14.r{n}.inPkg']

    return cols


# ============================================================
# Namuna yozuvlar (sample data)
# ============================================================

def sample_row_inbound_clean():
    """Toza Inbound public — yapon turist."""
    return {
        '#': 1, 'ID': 'a1b2c3d4-5678-90ab-cdef-1234567890ab',
        'Turi': 'Inbound (xorijliklar)', 'Manba': "Turist o'zi to'ldirgan", 'Til': 'en',
        'Boshlangan': '2026-05-15 14:23:11', 'Yakunlangan': '2026-05-15 14:28:32',
        "To'ldirish (s)": 321,
        'Xodim login': '', 'Xodim F.I.SH': '', 'Xodim telefon': '',
        "Bo'lim kodi": '', "Bo'lim nomi": '', 'Region kodi': '', 'Post turi': '',
        'IP': '203.0.113.42', 'Subnet /24': '203.0.113.0/24',
        'Device turi': 'mobile', 'OS': 'iOS', 'Brauzer': 'Safari',
        'Screen': '375x812', 'Platform': 'iPhone', 'Til (client)': 'ja-JP', 'Timezone': 'Asia/Tokyo',
        'User-Agent': 'Mozilla/5.0 (iPhone; CPU iPhone OS 17_4 like Mac OS X)...',
        'GPS bormi': 'Ha', 'Latitude': 41.3275, 'Longitude': 69.2817,
        'GPS aniqlik (m)': 18.5, 'Google Maps URL': 'https://www.google.com/maps?q=41.3275,69.2817',
        'Filtr holati': "O'tdi", 'F1': 'yes', 'F2': 'no', 'F3': '',
        'Davlat': 'Japan', 'Maqsad': 'leisure', 'Tunlar': 5,
        'Xarajat (umumiy)': 2500.00, 'Valyuta': 'USD',
        'IN.Q1 (Doimiy davlat)': 'Japan',
        'IN.Q2 (Pasport)': 'same',
        'IN.Q3 (Maqsad)': 'leisure',
        'IN.Q5_nights': 5, 'IN.Q5_zero': "Yo'q",
        'IN.Q6.Tashkent': 2, 'IN.Q6.Samarkand': 2, 'IN.Q6.Bukhara': 1,
        'IN.Q7 (Turar joy)': 0,
        'IN.Q8.restaurants': 'Ha', 'IN.Q8.national': 'Ha',
        'IN.Q9 (Paket?)': 'yes',
        'IN.Q10_total': 7, 'IN.Q10_uz': 5, 'IN.Q11 (Paket kishi)': 2,
        'IN.Q12_amount': 1800, 'IN.Q12_currency': 'USD',
        'IN.Q13 (Kelish)': 'airplane', 'IN.Q13_airline': 'uzair',
        'IN.Q14 (Ketish)': 'airplane', 'IN.Q14_airline': 'other',
        'IN.Q16_sum': 2500, 'IN.Q16_currency': 'USD', 'IN.Q16_persons': 2,
        'IN.Q17.r1.amount': 600, 'IN.Q17.r1.currency': 'USD', 'IN.Q17.r1.inPackage': "Yo'q",
        'IN.Q17.r2.amount': 350, 'IN.Q17.r2.currency': 'USD', 'IN.Q17.r2.inPackage': "Yo'q",
        'IN.Q17.r12.amount': 400, 'IN.Q17.r12.currency': 'USD', 'IN.Q17.r12.inPackage': "Yo'q",
        'IN.Q18.r0_intl_transport': 9, 'IN.Q18.r1_passport_control': 8,
        'IN.Q18.r2_hospitality': 10, 'IN.Q18.r3_value_for_money': 8,
        'IN.Q18.r4_food': 9, 'IN.Q18.r5_cleanliness': 8,
        'IN.Q18.r6_local_transport': 9, 'IN.Q18.r7_safety': 10,
        'IN.Q18.r8_culture': 10, 'IN.Q18.r9_accommodation': 9,
        'IN.Q19 (Izoh)': 'Beautiful country, excellent hospitality. Recommended!',
    }


def sample_row_outbound_clean():
    """Toza Outbound public — UZ fuqarosi Turkiyadan qaytgan."""
    return {
        '#': 2, 'ID': 'b2c3d4e5-6789-01bc-defa-2345678901bc',
        'Turi': 'Outbound (chiqish)', 'Manba': "Turist o'zi to'ldirgan", 'Til': 'uz',
        'Boshlangan': '2026-05-16 09:11:50', 'Yakunlangan': '2026-05-16 09:16:18',
        "To'ldirish (s)": 268,
        'Xodim login': '', 'Xodim F.I.SH': '', 'Xodim telefon': '',
        "Bo'lim kodi": '', "Bo'lim nomi": '', 'Region kodi': '', 'Post turi': '',
        'IP': '195.158.16.45', 'Subnet /24': '195.158.16.0/24',
        'Device turi': 'mobile', 'OS': 'Android', 'Brauzer': 'Chrome',
        'Screen': '412x915', 'Platform': 'Linux armv8l', 'Til (client)': 'uz-UZ', 'Timezone': 'Asia/Tashkent',
        'User-Agent': 'Mozilla/5.0 (Linux; Android 14; SM-S921B)...',
        'GPS bormi': 'Ha', 'Latitude': 41.2579, 'Longitude': 69.2812,
        'GPS aniqlik (m)': 22.3, 'Google Maps URL': 'https://www.google.com/maps?q=41.2579,69.2812',
        'Filtr holati': "O'tdi", 'F1': 'yes', 'F2': 'no', 'F3': '',
        'Davlat': 'Turkey', 'Maqsad': 'leisure', 'Tunlar': 8,
        'Xarajat (umumiy)': 1800.00, 'Valyuta': 'USD',
        'OUT.Q1 (Asosiy davlat)': 'Turkey', 'OUT.Q2 (Maqsad)': 'leisure',
        'OUT.Q4_val (Tunlar)': 8,
        'OUT.Q5 (Turar joy)': 'hotel', 'OUT.Q6 (Paket?)': 'yes',
        'OUT.Q7 (Paket tunlari)': 9, 'OUT.Q8 (Paket kishi)': 3,
        'OUT.Q9_amount': 1500, 'OUT.Q9_currency': 'USD',
        'OUT.Q10 (Chiqish)': 'airplane', 'OUT.Q10_airline': 'uzair',
        'OUT.Q11 (Qaytish)': 'airplane', 'OUT.Q11_airline': 'uzair',
        'OUT.Q13_amount': 1800, 'OUT.Q13_currency': 'USD', 'OUT.Q13_persons': 3,
        'OUT.Q14.r1.amount': 400, 'OUT.Q14.r1.currency': 'USD', 'OUT.Q14.r1.inPkg': 'Ha',
        'OUT.Q14.r2.amount': 280, 'OUT.Q14.r2.currency': 'USD', 'OUT.Q14.r2.inPkg': 'Ha',
        'OUT.Q14.r11.amount': 350, 'OUT.Q14.r11.currency': 'USD', 'OUT.Q14.r11.inPkg': "Yo'q",
    }


def sample_row_staff_inbound():
    """Staff kiritgan Inbound."""
    return {
        '#': 3, 'ID': 'c3d4e5f6-7890-12cd-efab-3456789012cd',
        'Turi': 'Inbound (xorijliklar)', 'Manba': 'Xodim kiritgan', 'Til': 'ru',
        'Boshlangan': '2026-05-17 11:42:30', 'Yakunlangan': '2026-05-17 11:45:18',
        "To'ldirish (s)": 168,
        'Xodim login': 'tajibekov_baxtiyar', 'Xodim F.I.SH': 'Tajibekov Baxtiyar Nietullaevich',
        'Xodim telefon': '+998901234567',
        "Bo'lim kodi": '117260', "Bo'lim nomi": "Toshkent xalqaro aeroporti",
        'Region kodi': '1726', 'Post turi': 'Xalqaro aeroport',
        'IP': '178.218.100.51', 'Subnet /24': '178.218.100.0/24',
        'Device turi': 'desktop', 'OS': 'Windows', 'Brauzer': 'Chrome',
        'Screen': '1920x1080', 'Platform': 'Win32', 'Til (client)': 'ru-RU', 'Timezone': 'Asia/Tashkent',
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)...',
        'GPS bormi': 'Ha', 'Latitude': 41.2579, 'Longitude': 69.2812,
        'GPS aniqlik (m)': 25.0, 'Google Maps URL': 'https://www.google.com/maps?q=41.2579,69.2812',
        'Filtr holati': "O'tdi", 'F1': 'yes', 'F2': 'no', 'F3': '',
        'Davlat': 'China', 'Maqsad': 'business', 'Tunlar': 3,
        'Xarajat (umumiy)': 3500.00, 'Valyuta': 'USD',
        'IN.Q1 (Doimiy davlat)': 'China', 'IN.Q2 (Pasport)': 'same',
        'IN.Q3 (Maqsad)': 'business', 'IN.Q4 (Biznes turi)': 'conference',
        'IN.Q5_nights': 3, 'IN.Q5_zero': "Yo'q",
        'IN.Q6.Tashkent': 3,
        'IN.Q7 (Turar joy)': 0,
        'IN.Q8.restaurants': 'Ha',
        'IN.Q9 (Paket?)': 'no',
        'IN.Q13 (Kelish)': 'airplane', 'IN.Q13_airline': 'other',
        'IN.Q14 (Ketish)': 'airplane', 'IN.Q14_airline': 'other',
        'IN.Q16_sum': 3500, 'IN.Q16_currency': 'USD', 'IN.Q16_persons': 1,
        'IN.Q17.r1.amount': 900, 'IN.Q17.r1.currency': 'USD', 'IN.Q17.r1.inPackage': "Yo'q",
        'IN.Q17.r2.amount': 250, 'IN.Q17.r2.currency': 'USD', 'IN.Q17.r2.inPackage': "Yo'q",
        'IN.Q18.r0_intl_transport': 8, 'IN.Q18.r2_hospitality': 9, 'IN.Q18.r4_food': 8,
    }


def sample_row_fraud_bot():
    """Bot UA — fraud."""
    return {
        '#': 4, 'ID': 'd4e5f6a7-8901-23de-fabc-4567890123de',
        'Turi': 'Inbound (xorijliklar)', 'Manba': "Turist o'zi to'ldirgan", 'Til': 'en',
        'Boshlangan': '2026-05-17 03:15:00', 'Yakunlangan': '2026-05-17 03:15:02',
        "To'ldirish (s)": 2,
        'IP': '185.220.101.50', 'Subnet /24': '185.220.101.0/24',
        'Device turi': 'bot', 'OS': '', 'Brauzer': 'Other',
        'Screen': '', 'Platform': '', 'Til (client)': '', 'Timezone': '',
        'User-Agent': 'curl/8.4.0',
        'GPS bormi': "Yo'q", 'Latitude': '', 'Longitude': '',
        'GPS aniqlik (m)': '', 'Google Maps URL': '',
        'Filtr holati': "O'tdi", 'F1': 'yes', 'F2': 'no', 'F3': '',
        'Davlat': 'United States', 'Maqsad': 'leisure', 'Tunlar': 3,
        'IN.Q1 (Doimiy davlat)': 'United States', 'IN.Q2 (Pasport)': 'same',
        'IN.Q3 (Maqsad)': 'leisure',
        'IN.Q5_nights': 3, 'IN.Q5_zero': "Yo'q",
    }


# ============================================================
# Yaratish
# ============================================================

def main():
    columns = build_columns()
    print(f"Jami {len(columns)} ta ustun")

    wb = Workbook()
    ws = wb.active
    ws.title = "Master sheet (Inbound+Outbound)"

    # Stillar
    header_font = Font(bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill(start_color='0F3A6E', end_color='0F3A6E', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    in_fill = PatternFill(start_color='1A5FA8', end_color='1A5FA8', fill_type='solid')
    out_fill = PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid')

    # Headerlarni yozish
    ws.append(columns)

    # Header'ni rang berish: IN. va OUT. prefiksiga qarab
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = header_align
        if col_name.startswith('IN.'):
            cell.fill = in_fill
        elif col_name.startswith('OUT.'):
            cell.fill = out_fill
        else:
            cell.fill = header_fill

    # Sample rowlarni yozish
    samples = [
        sample_row_inbound_clean(),
        sample_row_outbound_clean(),
        sample_row_staff_inbound(),
        sample_row_fraud_bot(),
    ]
    for sample in samples:
        row_data = [sample.get(col, '') for col in columns]
        ws.append(row_data)

    # Column width — auto-fit (ehtiyot uchun cheklash)
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        width = max(len(col_name) + 2, 12)
        # IN.Q* / OUT.Q* shortroq
        if col_name.startswith(('IN.Q17.', 'IN.Q18.', 'OUT.Q14.')):
            width = max(width, 16)
        elif col_name.startswith(('IN.', 'OUT.')):
            width = max(width, 18)
        ws.column_dimensions[col_letter].width = min(width, 40)

    # Freeze: 1-qator + 1-8 ustunlar (ID/Turi/Manba)
    ws.freeze_panes = 'I2'

    # Row height
    ws.row_dimensions[1].height = 45

    # Saqlash
    output_path = Path(__file__).parent / 'master_sheet_sample.xlsx'
    wb.save(output_path)
    print(f"[OK] Saqlandi: {output_path}")
    print(f"   Ustunlar soni: {len(columns)}")
    print(f"   Sample qatorlar: {len(samples)}")
    print()
    print("Ustunlar guruhi:")
    print("  1-8   Identifikatsiya va meta")
    print("  9-15  Xodim va bo'lim")
    print("  16-25 Tarmoq va device")
    print("  26-30 GPS")
    print("  31-34 Screening")
    print("  35-39 Auto-summary")
    print(f"  40-{40 + len([c for c in columns if c.startswith('IN.')]) - 1}  Inbound (Q1-Q19)")
    out_start = 40 + len([c for c in columns if c.startswith('IN.')])
    print(f"  {out_start}-{len(columns)}  Outbound (Q1-Q14)")


if __name__ == '__main__':
    main()
